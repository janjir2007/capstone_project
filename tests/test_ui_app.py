"""Screen tests for the Streamlit app (tasks 5.1–5.3c, DESIGN v2) with a FAKE backend.

The app runs in Streamlit's AppTest (no browser). The network is replaced by
httpx.MockTransport + FakeApi, a tiny in-memory imitation of the endpoints in
docs/api_contract.md. No real API, no database, invented data only (RISK-06).

Skipped when Streamlit is not installed (CI installs only requirements.txt;
install requirements-ui.txt to run these).
"""

import json
from typing import Any
from urllib.parse import unquote
from pathlib import Path

import httpx
import pytest

st = pytest.importorskip("streamlit")
from streamlit.testing.v1 import AppTest  # noqa: E402

from src.ui import api_client  # noqa: E402

APP = str(Path(__file__).resolve().parent.parent / "src" / "ui" / "app.py")
TODAY = "2026-10-12"

DESTINATIONS = [
    {"id": 1, "name": "Ulaanbaatar city (local work)", "route_category": "ub_city"},
    {"id": 2, "name": "Zamyn-Uud (MN–CN border)", "route_category": "mn_china", "name_mn": "Замын-Үүд"},
    {"id": 5, "name": "Altanbulag (MN–RU border)", "route_category": "mn_russia", "name_mn": None},
    {"id": 6, "name": "Kyakhta (RU side of border)", "route_category": "mn_russia", "name_mn": ""},
    {"id": 7, "name": "Other / several / unknown", "route_category": "mixed"},
]


class FakeApi:
    """In-memory imitation of the data endpoints (only what the UI uses)."""

    def __init__(self) -> None:
        self.vehicles: list[dict] = []
        self.services: list[dict] = []
        self.posts: list[tuple[dict, bool]] = []  # (body, force)
        self.prediction: dict | None = None        # returned on every card (None = model not loaded)
        self.recognize: tuple[int, dict] = (503, {"detail": "plate recognizer not loaded"})
        self.uploads: list[bytes] = []              # raw multipart bodies sent to /plates/recognize
        self.reminders: list[dict] | None = []      # None -> 503 (model not loaded)
        self.sent: dict[int, dict] = {}             # service_id -> stored "sent" row
        self.hide_sent = True                       # False = stale list (another manager's open tab)
        self.destinations: list[dict] = DESTINATIONS  # GET /destinations
        self.searches: list[str] = []               # q of every GET /vehicles/search (9.1)
        self.search_params: list[dict] = []         # all params of every GET /vehicles/search (9.1 step 2)

    # -- data ---------------------------------------------------------------
    def add_vehicle(self, identifier: str, company: str, identifier_type: str = "plate", **service) -> dict:
        vehicle = {
            "id": len(self.vehicles) + 1,
            "identifier_type": identifier_type,
            "identifier": identifier,
            "vehicle_model": service.pop("vehicle_model", None),
            "company": {"id": len(self.vehicles) + 1, "name": company},
        }
        self.vehicles.append(vehicle)
        if service:
            self._add_service(vehicle["id"], service)
        return vehicle

    def _add_service(self, vehicle_id: int, body: dict) -> dict:
        dest = next((d for d in DESTINATIONS if d["id"] == body.get("destination_id")), None)
        service = {
            "id": len(self.services) + 1,
            "vehicle_id": vehicle_id,
            "service_date": body.get("service_date", TODAY),
            "odometer_km": body.get("odometer_km"),
            # 9.5a: the API adds odometer + OIL_INTERVAL_KM (default 15000); none without odometer
            "target_km": None if body.get("odometer_km") is None else body["odometer_km"] + 15000,
            "engine_hours": body.get("engine_hours"),
            "oil_grade": body["oil_grade"],
            "liters": body["liters"],
            "route": dest["route_category"] if dest else body.get("route"),
            "destination": dest,
            "destination_raw": body.get("destination_raw"),
            "driver_name": body.get("driver_name"),
            "driver_phone": body.get("driver_phone"),
            "oil_changer": body.get("oil_changer"),
            "photo_ref": None,
            "created_at": f"{TODAY}T10:00:00+08:00",
        }
        self.services.append(service)
        return service

    @staticmethod
    def _key(text: str) -> str:
        return "".join(text.split()).upper()

    def card(self, vehicle: dict) -> dict:
        history = sorted(
            (s for s in self.services if s["vehicle_id"] == vehicle["id"]),
            key=lambda s: (s["service_date"], s["id"]),
            reverse=True,
        )
        n = len(history)
        return {
            **vehicle,
            "plate_normalized": self._key(vehicle["identifier"]) if vehicle["identifier_type"] == "plate" else None,
            "service_count": n,
            "segment": None if n == 0 else "new" if n == 1 else "returning" if n == 2 else "loyal",
            "usual_liters": history[0]["liters"] if history else None,
            "last_service": history[0] if history else None,
            "history": history,
            "prediction": {**self.prediction, "vehicle_id": vehicle["id"]} if self.prediction and n else None,
        }

    # -- endpoints ----------------------------------------------------------
    def handle(self, request: httpx.Request) -> httpx.Response:
        path = request.url.path
        if path == "/destinations":
            return httpx.Response(200, json=self.destinations)
        if path.startswith("/vehicles/") and path.split("/")[-1].isdigit():
            vid = int(path.split("/")[-1])
            vehicle = next((v for v in self.vehicles if v["id"] == vid), None)
            return httpx.Response(200, json=self.card(vehicle)) if vehicle else httpx.Response(404, json={"detail": "unknown"})
        if path == "/stats/segments":
            return httpx.Response(200, json={"as_of": TODAY, "total": 7,
                                             "counts": {"new": 2, "returning": 3, "loyal": 1, "at_risk": 1}})
        if path == "/reminders/due":
            if self.reminders is None:
                return httpx.Response(503, json={"detail": "prediction model not loaded"})
            due = [r for r in self.reminders if not (self.hide_sent and r["service_id"] in self.sent)]
            return httpx.Response(200, json={"as_of": TODAY, "contact_is_placeholder": True, "reminders": due})
        if path.startswith("/reminders/") and path.endswith("/sent"):
            sid = int(path.split("/")[2])
            if sid == 999:
                return httpx.Response(409, json={"detail": {"code": "not_latest_service", "message": "newer",
                                                            "latest_service_id": 1000}})
            if sid in self.sent:
                return httpx.Response(200, json={**self.sent[sid], "already_sent": True})
            self.sent[sid] = {"id": len(self.sent) + 1, "service_id": sid, "phone": "99110002",
                              "sent_at": f"{TODAY}T11:00:00+08:00"}
            return httpx.Response(201, json={**self.sent[sid], "already_sent": False})
        if path == "/vehicles/search":
            self.search_params.append(dict(request.url.params))
            return self._search(request.url.params.get("q", ""), int(request.url.params.get("limit", 200)),
                                request.url.params.get("company_id"), request.url.params.get("with_services") == "true")
        if path == "/vehicles/lookup":
            return self._lookup(request.url.params["q"])
        if path == "/plates/recognize" and request.method == "POST":
            self.uploads.append(request.read())
            status, body = self.recognize
            return httpx.Response(status, json=body)
        if path == "/services" and request.method == "POST":
            return self._post(json.loads(request.content), request.url.params.get("force") == "true")
        return httpx.Response(404, json={"detail": "not found"})

    def _search(self, q: str, limit: int, company_id: str | None = None, with_services: bool = False) -> httpx.Response:
        """GET /vehicles/search: plate or company contains q, + company filter (enough for the screen
        tests; the other filters are only checked as sent params). Options from every truck."""
        self.searches.append(q)
        key = self._key(q)
        rows = []
        for v in self.vehicles:
            card = self.card(v)
            if key and key not in self._key(v["identifier"]) and key not in self._key(v["company"]["name"]):
                continue
            if company_id is not None and v["company"]["id"] != int(company_id):
                continue
            rows.append({k: card[k] for k in ("id", "identifier_type", "identifier", "vehicle_model", "company",
                                              "service_count", "last_service", "segment")}
                        | {"status": None, "days_left": None, "matched_on": [],
                           "services": card["history"] if with_services else None})
        options = {"companies": sorted((v["company"] for v in self.vehicles), key=lambda c: c["name"]),
                   "oil_grades": sorted({s["oil_grade"] for s in self.services})}
        return httpx.Response(200, json={"query": q, "total": len(rows), "vehicles": rows[:limit], "options": options})

    def _lookup(self, q: str) -> httpx.Response:
        digits = "".join(c for c in q if c.isdigit())
        if digits.startswith("976") and len(digits) == 11:
            digits = digits[3:]
        if len(digits) == 8 and not any(c.isalpha() for c in q):
            ids = {s["vehicle_id"] for s in self.services if s["driver_phone"] == digits}
            found = [v for v in self.vehicles if v["id"] in ids]
            return httpx.Response(200, json={
                "query": q, "match_type": "phone", "normalized": digits,
                "vehicles": [self.card(v) for v in found],
            })
        found = [v for v in self.vehicles if self._key(v["identifier"]) == self._key(q)]
        match_type = "fleet_no" if found and all(v["identifier_type"] == "fleet_no" for v in found) else "plate"
        return httpx.Response(200, json={
            "query": q, "match_type": match_type, "normalized": self._key(q),
            "vehicles": [self.card(v) for v in found],
        })

    def _post(self, body: dict, force: bool) -> httpx.Response:
        self.posts.append((body, force))
        kind = body.get("identifier_type", "plate")
        vehicle = next(
            (v for v in self.vehicles
             if v["identifier_type"] == kind and self._key(v["identifier"]) == self._key(body["identifier"])
             and (kind == "plate" or v["company"]["name"] == body.get("company_name"))),
            None,
        )
        created = vehicle is None
        if created:
            if not body.get("company_name"):
                return httpx.Response(422, json={"detail": "new vehicle: company_name is required (private owner = owner's name)"})
            vehicle = self.add_vehicle(body["identifier"], body["company_name"], kind,
                                       vehicle_model=body.get("vehicle_model"))
        elif not force and any(s["vehicle_id"] == vehicle["id"] and s["service_date"] == TODAY for s in self.services):
            existing = next(s for s in self.services if s["vehicle_id"] == vehicle["id"] and s["service_date"] == TODAY)
            return httpx.Response(409, json={"detail": {
                "code": "duplicate_service", "message": "already has a service on this date",
                "existing_service_id": existing["id"],
            }})
        service = self._add_service(vehicle["id"], body)
        return httpx.Response(201, json={
            "service": service, "vehicle_id": vehicle["id"],
            "vehicle_created": created, "company_created": created,
        })


@pytest.fixture
def fake_api(monkeypatch):
    """Fake backend behind the UI's HTTP client; login off; fresh cache."""
    monkeypatch.delenv("UI_PASSWORD", raising=False)
    fake = FakeApi()
    api_client.set_transport(httpx.MockTransport(fake.handle))
    st.cache_data.clear()
    yield fake
    api_client.set_transport(None)


HOME, RECORD, SEARCH, CARD, REMINDERS, DASHBOARD = (
    "views/home.py", "views/record.py", "views/search.py", "views/card.py", "views/reminders.py", "views/dashboard.py")

PREDICTION = {
    "method": "ridge", "fallback_reason": None, "model_version": "ridge-0000abcd",
    "based_on_service_id": 1, "last_service_date": "2026-09-01", "predicted_days": 50.0,
    "predicted_date": "2026-10-21", "today": TODAY, "days_left": 9, "status": "due_soon",
}
RECOGNIZED = {"plate_raw": "0002ТЕС", "plate_normalized": "0002ТЕС", "confidence": 0.95,
              "needs_manual": False, "reason": None, "detection": None,
              "detector": "yolo26n (onnx)", "recognizer": "char_cnn (onnx)", "ms": 80.0}


def run_app(page: str | None = None) -> AppTest:
    at = AppTest.from_file(APP, default_timeout=10).run()
    return at.switch_page(page).run() if page else at


def on(at: AppTest, page: str) -> AppTest:
    """AppTest does not remember an in-app st.switch_page for the NEXT run
    (a browser does, via the URL): select the page again before interacting."""
    return at.switch_page(page)


def click(at: AppTest, label: str) -> AppTest:
    """Click the button with this label and rerun."""
    matches = [b for b in at.button if b.label == label]
    assert matches, f"no button {label!r}; have {[b.label for b in at.button]}"
    return matches[0].click().run()


def texts(elements) -> list[str]:
    return [e.value for e in elements]


def html_blocks(at: AppTest) -> str:
    return "\n".join(texts(at.markdown))


def step(at: AppTest) -> str:
    return at.session_state["rec_step"]


def type_plate(at: AppTest, plate: str, kind: str = "plate") -> AppTest:
    """Record start: «Гараар бичих» (or «Дугааргүй машин») -> type -> «Үргэлжлүүлэх» -> form
    (a fleet number found in a company -> confirm)."""
    at = click(at, "Дугааргүй машин" if kind == "fleet_no" else "Гараар бичих")
    at.text_input(key="rec_manual").set_value(plate).run()
    return click(at, "Үргэлжлүүлэх")


def to_form(at: AppTest, plate: str, kind: str = "plate") -> AppTest:
    at = type_plate(at, plate, kind)
    return click(at, "Зөв, үргэлжлүүлэх") if step(at) == "confirm" else at


def fill_new_vehicle(at: AppTest, company: str | None = None) -> AppTest:
    at.number_input(key="f_odometer_km").set_value(120000)
    if company is not None:  # the field exists only in fleet-number mode (T8)
        at.text_input(key="f_company_name").set_value(company)
    at.selectbox(key="f_model_choice").set_value("Foton Auman")
    at.number_input(key="f_liters").set_value(28.5)
    at.text_input(key="f_driver_name").set_value("Тест Б.")
    at.text_input(key="f_driver_phone").set_value("99110099")
    return at.run()


def seed_existing(fake: FakeApi) -> dict:
    return fake.add_vehicle(
        "0002 ТЕС", "Говь Тест ХХК", vehicle_model="Shacman X3000",
        service_date="2026-09-01", odometer_km=200000, oil_grade="10W-40", liters=32.0,
        destination_id=2, driver_name="Жолооч Д.", driver_phone="99110002",
    )


def upload_photo(at: AppTest, key: str) -> AppTest:
    return at.file_uploader(key=key).upload("truck.jpg", b"fake-jpeg", "image/jpeg").run()


@pytest.fixture
def phone(monkeypatch):
    """Fake plate camera (task 8.4). The browser part cannot run in AppTest, so the record
    screen's camera() / gallery_button() are replaced: shoot(at, ...) makes the NEXT run get one
    photo, exactly like the component's one-shot trigger value."""
    from src.ui import native_camera
    pending: dict[str, native_camera.CapturedPhoto] = {}

    def mount(source: str):
        return lambda key, max_side=native_camera.MAX_SIDE: pending.pop(source, None)

    monkeypatch.setattr(native_camera, "camera", mount("live"))
    monkeypatch.setattr(native_camera, "gallery_button", mount("gallery"))

    def shoot(at: AppTest, data: bytes = b"fake-jpeg", name: str = "photo.jpg", type_: str = "image/jpeg",
              source: str = "live", error: str | None = None) -> AppTest:
        """error="..." -> the browser could not read the photo (fail() in the component, 8.4b)."""
        pending[source] = (native_camera.PhotoFailed(error, source) if error else
                           native_camera.CapturedPhoto(data, name, type_, 1920, 2560, None, False, source, 1920, 2560))
        return at.run()

    return shoot


def tiny_jpeg() -> bytes:
    import io
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (40, 20), (200, 200, 200)).save(buf, format="JPEG")
    return buf.getvalue()


# --- login (D-061) + home ------------------------------------------------------------

def test_no_password_env_opens_home(fake_api) -> None:
    at = run_app()
    assert not at.exception
    assert texts(at.title) == ["Сайн байна уу"]
    assert {"Үйлчилгээ бүртгэх", "Хайх"} <= {b.label for b in at.button}
    assert "Гео-Эксперт" in html_blocks(at)


def test_login_required_and_works(fake_api, monkeypatch) -> None:
    monkeypatch.setenv("UI_PASSWORD", "demo-pass")
    at = run_app()
    assert texts(at.title) == ["Нэвтрэх"]
    at.text_input[0].set_value("wrong")
    at = click(at, "Нэвтрэх")
    assert "Нууц үг буруу байна." in texts(at.error)
    at.text_input[0].set_value("demo-pass")
    at = click(at, "Нэвтрэх")
    assert texts(at.title) == ["Сайн байна уу"]


ALL_VIEWS = [HOME, RECORD, SEARCH, CARD, REMINDERS, DASHBOARD]


@pytest.mark.parametrize("view", ALL_VIEWS)
def test_logged_out_cannot_open_any_view(fake_api, monkeypatch, view) -> None:
    """Login-bypass guard: with UI_PASSWORD set and not logged in, no screen is a
    navigation page at all, so a direct URL only gets the login screen."""
    monkeypatch.setenv("UI_PASSWORD", "demo-pass")
    at = run_app()
    with pytest.raises(ValueError, match="Could not find a navigation page"):
        at.switch_page(view)
    assert texts(at.title) == ["Нэвтрэх"]


@pytest.mark.parametrize("view", ALL_VIEWS)
def test_logged_in_can_open_every_view(fake_api, view) -> None:
    """Counter-check for the test above: the same views ARE pages once allowed."""
    at = run_app(view)
    assert not at.exception
    assert texts(at.title) and texts(at.title) != ["Нэвтрэх"]


def test_home_buttons_open_screens(fake_api) -> None:
    at = click(run_app(), "Үйлчилгээ бүртгэх")
    assert texts(at.title) == ["Үйлчилгээ бүртгэх"] and step(at) == "capture"
    at = click(run_app(), "Хайх")
    assert texts(at.title) == ["Хайх"]


@pytest.mark.parametrize("view", [SEARCH, DASHBOARD])
def test_back_button_goes_home(fake_api, view) -> None:
    """8.1: the manager screen's back arrow also goes home (was «Сануулга»)."""
    at = click(run_app(view), "Буцах")
    assert texts(at.title) == ["Сайн байна уу"]


# --- record flow (5.2 / DESIGN v2) ------------------------------------------------------

def test_new_vehicle_full_flow(fake_api) -> None:
    at = type_plate(run_app(RECORD), "0001 ТЕС")
    assert step(at) == "form"                                    # T9: no separate confirm step for a plate
    assert texts(at.title) == ["Үйлчилгээ бүртгэх"]
    assert '<div class="vh-id">0001 ТЕС</div>' in html_blocks(at)
    assert "Шинэ машин" in html_blocks(at)
    at = click(fill_new_vehicle(at), "Хадгалах")
    assert not at.exception, at.exception
    body, force = fake_api.posts[-1]
    assert force is False
    assert body == {
        "identifier_type": "plate", "identifier": "0001 ТЕС", "oil_grade": "15W-40", "liters": 28.5,
        "company_name": "Тодорхойгүй", "vehicle_model": "Foton Auman", "odometer_km": 120000,
        "driver_name": "Тест Б.", "driver_phone": "99110099",
    }  # 9.5a: no destination / route
    assert step(at) == "success"
    assert "Хадгалагдлаа" in html_blocks(at)
    assert {"Дараагийн машин", "Машины карт үзэх"} <= {b.label for b in at.button}

    at = click(at, "Дараагийн машин")
    assert step(at) == "capture"
    assert not any(w.label == "Тос солигч" for w in at.text_input)  # field removed (T6)


def test_existing_vehicle_confirm_chip_and_prefill(fake_api) -> None:
    seed_existing(fake_api)
    at = type_plate(run_app(RECORD), "0002тес")
    assert step(at) == "form" and texts(at.title) == ["Үйлчилгээ бүртгэх"]  # typed -> full form (mockup p3)
    assert "Бүртгэлтэй · 1 удаа ирсэн" in html_blocks(at)
    state = at.session_state
    assert (state["f_oil_choice"], state["f_liters"], state["f_driver_phone"]) == ("10W-40", 32.0, "99110002")
    assert state["f_odometer_km"] is None and "f_destination_id" not in state  # fresh; 9.5a: no destination
    assert not at.pills  # 9.5a: no destination pills
    assert texts(at.subheader) == ["Энэ удаа", "Жолооч"]
    assert "Өмнөх үйлчилгээнээс бөглөгдсөн" in html_blocks(at)
    assert "Өмнөх: 200 000 км" in html_blocks(at)
    assert {"Засах", "Хадгалах"} <= {b.label for b in at.button}


def test_new_vehicle_defaults_and_hidden_fields(fake_api) -> None:
    at = to_form(run_app(RECORD), "0001 ТЕС")
    assert at.session_state["f_oil_choice"] == "15W-40"                # default oil grade
    assert at.segmented_control(key="f_oil_choice").options == ["15W-40", "10W-40", "20W-50", "Бусад"]
    assert "Компани / эзэмшигч" not in [w.label for w in at.text_input]  # plate: no company field (T8)
    assert "Volvo FH" in at.selectbox(key="f_model_choice").options


def test_new_fleet_vehicle_asks_company_and_sends_it(fake_api) -> None:
    """T8: «Дугааргүй машин» mode keeps the company field (fleet numbers are unique per company)."""
    at = to_form(run_app(RECORD), "HT-99", kind="fleet_no")
    assert "Компани / эзэмшигч" in [w.label for w in at.text_input]
    at.number_input(key="f_engine_hours").set_value(500.0)
    at = click(fill_new_vehicle(at, company="Уурхай В ХХК"), "Хадгалах")
    body, _ = fake_api.posts[-1]
    assert (body["identifier_type"], body["company_name"]) == ("fleet_no", "Уурхай В ХХК")


def test_new_fleet_vehicle_without_company_is_blocked(fake_api) -> None:
    at = to_form(run_app(RECORD), "HT-99", kind="fleet_no")
    at = click(fill_new_vehicle(at, company=""), "Хадгалах")
    assert "Компанийн нэр (эсвэл эзэмшигч) оруулна уу." in texts(at.error)
    assert fake_api.posts == []


def test_existing_vehicle_hides_company_and_model(fake_api) -> None:
    seed_existing(fake_api)
    at = to_form(run_app(RECORD), "0002 ТЕС")
    labels = [w.label for w in at.text_input]
    assert "Компани / эзэмшигч" not in labels and not at.selectbox
    at.number_input(key="f_odometer_km").set_value(215000).run()
    at = click(at, "Хадгалах")
    body, _ = fake_api.posts[-1]
    assert "company_name" not in body and "vehicle_model" not in body and body["oil_grade"] == "10W-40"


def test_other_oil_grade_and_other_model_saved_as_typed(fake_api) -> None:
    at = to_form(run_app(RECORD), "0001 ТЕС")
    at.number_input(key="f_odometer_km").set_value(1000)
    at.number_input(key="f_liters").set_value(30.0)
    at.segmented_control(key="f_oil_choice").set_value("Бусад")
    at.selectbox(key="f_model_choice").set_value("Бусад").run()
    at.text_input(key="f_oil_other").set_value("5W-30")
    at.text_input(key="f_model_other").set_value("Sitrak C7H").run()
    at = click(at, "Хадгалах")
    body, _ = fake_api.posts[-1]
    assert (body["oil_grade"], body["vehicle_model"]) == ("5W-30", "Sitrak C7H")


def test_existing_fleet_truck_sends_its_company(fake_api) -> None:
    """Company is hidden for existing trucks, but a fleet number still needs it (D-042)."""
    fake_api.add_vehicle("HT-12", "Уурхай А ХХК", "fleet_no", oil_grade="15W-40", liters=40.0, engine_hours=1000.0)
    at = to_form(run_app(RECORD), "HT-12", kind="fleet_no")
    at.number_input(key="f_engine_hours").set_value(1100.0).run()
    at = click(at, "Хадгалах")
    body, _ = fake_api.posts[-1]
    assert body["company_name"] == "Уурхай А ХХК" and body["identifier_type"] == "fleet_no"


def test_photo_preview_on_result_and_form(fake_api, phone) -> None:
    seed_existing(fake_api)
    fake_api.recognize = (200, RECOGNIZED)
    at = phone(run_app(RECORD), tiny_jpeg())
    assert step(at) == "result" and '<div class="rec-photo"><img' in html_blocks(at)  # 8.4 result screen
    at = click(at, "Үргэлжлүүлэх")
    assert step(at) == "form" and texts(at.title) == ["Машин танигдлаа"]
    assert 'src="data:image/jpeg;base64,' in html_blocks(at)        # plate strip from the photo (T9)
    assert "Автоматаар уншсан" in html_blocks(at)
    assert len(fake_api.uploads) == 1                               # sent once, to /plates/recognize only
    at = click(at, "Буцах")                                        # leave; photo stays session-only


def test_unreadable_photo_shows_text_not_crash(fake_api, phone) -> None:
    fake_api.recognize = (200, RECOGNIZED)
    at = phone(run_app(RECORD), b"not-decodable", "IMG_0002.HEIC", "", source="gallery")
    assert not at.exception
    assert "урьдчилж харах боломжгүй" in html_blocks(at)


def test_wrong_plate_goes_back_to_edit(fake_api) -> None:
    at = click(type_plate(run_app(RECORD), "0001 ТЕС"), "Засах")
    assert step(at) == "result"
    assert at.text_input(key="rec_manual").value == "0001 ТЕС"


def test_empty_form_shows_mongolian_errors_and_sends_nothing(fake_api) -> None:
    at = click(to_form(run_app(RECORD), "0001 ТЕС"), "Хадгалах")
    assert texts(at.error)                                                 # Mongolian messages
    assert "Компанийн нэр (эсвэл эзэмшигч) оруулна уу." not in texts(at.error)  # plate: company automatic (T8)
    assert fake_api.posts == []


def test_odometer_lower_than_last_warns_then_saves_on_confirm(fake_api) -> None:
    """Task 9.5a: a new odometer not higher than the last one -> note under the field, and
    «Хадгалах» asks first («Гүйлт зөв үү?»); «Ингээд хадгалах» saves it anyway."""
    seed_existing(fake_api)  # last odometer 200 000 km
    at = to_form(run_app(RECORD), "0002 ТЕС")
    at.number_input(key="f_odometer_km").set_value(20000).run()  # a digit missing
    assert "Гүйлт өмнөхөөсөө (200 000 км) их байх ёстой" in html_blocks(at)
    at = click(at, "Хадгалах")
    assert "Гүйлт зөв үү? Өмнөх: 200 000 км · одоо: 20 000 км" in texts(at.warning)
    assert fake_api.posts == []                                    # nothing sent yet
    at = click(at, "Ингээд хадгалах")
    assert step(at) == "success" and fake_api.posts[-1][0]["odometer_km"] == 20000
    assert fake_api.posts[-1][1] is False                          # not a duplicate override


def test_odometer_warning_fix_and_save(fake_api) -> None:
    """«Гүйлтээ засах» closes the question; a corrected (higher) odometer saves with no question."""
    seed_existing(fake_api)
    at = to_form(run_app(RECORD), "0002 ТЕС")
    at.number_input(key="f_odometer_km").set_value(200000).run()   # equal is not higher
    at = click(click(at, "Хадгалах"), "Гүйлтээ засах")
    assert not texts(at.warning) and fake_api.posts == []
    at.number_input(key="f_odometer_km").set_value(201500).run()
    assert "их байх ёстой" not in html_blocks(at)
    at = click(at, "Хадгалах")
    assert step(at) == "success" and fake_api.posts[-1][0]["odometer_km"] == 201500


def test_no_odometer_warning_for_new_or_mine_truck(fake_api) -> None:
    """Nothing to compare: a new truck, or a mine truck with engine hours only."""
    fake_api.add_vehicle("HT-12", "Уурхай А ХХК", "fleet_no", service_date="2026-09-01", oil_grade="15W-40",
                         liters=40.0, engine_hours=1000.0)
    at = to_form(run_app(RECORD), "HT-12", kind="fleet_no")
    at.number_input(key="f_engine_hours").set_value(1100.0)
    at.number_input(key="f_odometer_km").set_value(5).run()
    at = click(at, "Хадгалах")
    assert step(at) == "success"


def test_duplicate_then_force_save(fake_api) -> None:
    vehicle = seed_existing(fake_api)
    fake_api._add_service(vehicle["id"], {"oil_grade": "10W-40", "liters": 32.0, "odometer_km": 210000})
    at = to_form(run_app(RECORD), "0002 ТЕС")
    at.number_input(key="f_odometer_km").set_value(210050).run()
    at = click(at, "Хадгалах")
    assert any("аль хэдийн бүртгэгдсэн" in t for t in texts(at.warning))
    count = len(fake_api.services)
    at = click(at, "Ингээд хадгалах")
    assert fake_api.posts[-1][1] is True
    assert len(fake_api.services) == count + 1
    assert step(at) == "success"


def test_duplicate_cancel(fake_api) -> None:
    vehicle = seed_existing(fake_api)
    fake_api._add_service(vehicle["id"], {"oil_grade": "10W-40", "liters": 32.0, "odometer_km": 210000})
    at = to_form(run_app(RECORD), "0002 ТЕС")
    at.number_input(key="f_odometer_km").set_value(210050).run()
    at = click(click(at, "Хадгалах"), "Болих")
    assert "Хадгалах" in [b.label for b in at.button]
    assert len(fake_api.posts) == 1


def test_fleet_number_in_two_companies(fake_api) -> None:
    fake_api.add_vehicle("HT-12", "Уурхай А ХХК", "fleet_no", oil_grade="15W-40", liters=40.0, engine_hours=1000.0)
    fake_api.add_vehicle("HT-12", "Уурхай Б ХХК", "fleet_no", oil_grade="10W-30", liters=45.0, engine_hours=2000.0)
    at = type_plate(run_app(RECORD), "ht-12", kind="fleet_no")
    at.radio(key="rec_choice").set_value(1).run()
    at = click(at, "Зөв, үргэлжлүүлэх")
    assert at.session_state["f_company_name"] == "Уурхай Б ХХК"
    # 10W-30 is not one of the 3 tap options -> «Бусад» + the old value as text
    assert (at.session_state["f_oil_choice"], at.session_state["f_oil_other"]) == ("Бусад", "10W-30")
    assert at.number_input(key="f_engine_hours")  # mine trucks: engine hours shown


def test_api_down_shows_message(fake_api) -> None:
    def down(request: httpx.Request) -> httpx.Response:
        raise httpx.ConnectError("refused", request=request)

    api_client.set_transport(httpx.MockTransport(down))
    at = type_plate(run_app(RECORD), "0001 ТЕС")
    assert not at.exception
    assert any("холбогдож" in t for t in texts(at.error))
    assert step(at) == "result"


# --- photo -> recognize ------------------------------------------------------------------

def test_photo_recognized_existing_truck(fake_api, phone) -> None:
    """Photo -> 8.4 result (green, plate pre-filled) -> «Үргэлжлүүлэх» -> T9 «Машин танигдлаа»:
    truck line, previous visit + target km (9.5a, instead of the arrival block), oil/driver rows."""
    seed_existing(fake_api)
    fake_api.prediction = PREDICTION
    fake_api.recognize = (200, {**RECOGNIZED, "detection": {"box": [1, 1, 30, 10], "confidence": 0.9}})
    at = phone(run_app(RECORD), tiny_jpeg())
    assert not at.exception
    assert b'name="file"' in fake_api.uploads[0]
    assert step(at) == "result" and texts(at.title) == ["Үйлчилгээ бүртгэх"]
    assert "Дугаарыг таньлаа — шалгаад зөв бол үргэлжлүүлнэ үү" in html_blocks(at)
    assert at.text_input(key="rec_manual").value == "0002ТЕС"
    assert len(fake_api.uploads) == 1 and "rec_checked" not in at.session_state  # no auto lookup (8.4)
    at = click(at, "Үргэлжлүүлэх")
    assert step(at) == "form" and texts(at.title) == ["Машин танигдлаа"]
    page = html_blocks(at)
    assert "Шинэ · 1 удаа · Говь Тест ХХК" in page
    assert "Цагтаа ирсэн" not in page and "Таамаглаж байсан огноо" not in page  # 9.5a: no predicted date
    assert "09-01 · 200 000 км · Замын-Үүд" in page                 # «Өмнөх удаа» (old record keeps its place)
    assert "Дараагийн солилт" in page and "215 000 км" in page     # 9.5a: 200 000 + 15 000
    assert "10W-40 · 32 л" in page and "Жолооч Д. · 9911 0002" in page
    assert at.expander[0].label == "Тос, жолооч өөрчлөх"
    assert not at.pills                                            # 9.5a: no destination question
    at.number_input(key="f_odometer_km").set_value(215000)
    at = click(at, "Хадгалах")
    assert step(at) == "success" and "destination_id" not in fake_api.posts[-1][0]
    assert "company_name" not in fake_api.posts[-1][0]             # existing plate: no company (T8)


def test_recognized_wrong_goes_back_to_edit(fake_api, phone) -> None:
    seed_existing(fake_api)
    fake_api.recognize = (200, RECOGNIZED)
    at = click(click(phone(run_app(RECORD)), "Үргэлжлүүлэх"), "Буруу")
    assert step(at) == "result" and at.text_input(key="rec_manual").value == "0002ТЕС"
    assert "таньлаа" not in html_blocks(at)  # no green banner once he said it is wrong


def test_photo_needs_manual_shows_orange_and_empty_field(fake_api, phone) -> None:
    """8.4: not read -> orange message + EMPTY field (a partial read is not pre-filled)."""
    fake_api.recognize = (200, {**RECOGNIZED, "plate_normalized": None, "plate_raw": "0002ТЕ",
                                "needs_manual": True, "reason": "invalid_format"})
    at = phone(run_app(RECORD))
    assert step(at) == "result"
    assert at.text_input(key="rec_manual").value == ""
    page = html_blocks(at)
    assert "Дугаарыг уншиж чадсангүй — гараар оруулна уу" in page and "4 тоо + 3 үсэг" in page
    assert {"Үргэлжлүүлэх", "Дахин зураг авах"} <= {b.label for b in at.button}


def test_edited_plate_opens_normal_form(fake_api, phone) -> None:
    """A read plate the oil changer corrected is no longer «Машин танигдлаа» (not auto-read)."""
    seed_existing(fake_api)
    fake_api.add_vehicle("0003 ТЕС", "Тест ХХК", oil_grade="15W-40", liters=30.0)
    fake_api.recognize = (200, RECOGNIZED)
    at = phone(run_app(RECORD))
    at.text_input(key="rec_manual").set_value("0003 ТЕС").run()
    at = click(at, "Үргэлжлүүлэх")
    assert step(at) == "form" and texts(at.title) == ["Үйлчилгээ бүртгэх"]
    assert at.session_state["rec_checked"]["identifier"] == "0003 ТЕС"


def test_retake_returns_to_camera_and_forgets_photo(fake_api, phone) -> None:
    fake_api.recognize = (200, RECOGNIZED)
    at = click(phone(run_app(RECORD), tiny_jpeg()), "Дахин зураг авах")
    assert step(at) == "capture"
    assert "rec_photo" not in at.session_state and "rec_read" not in at.session_state


def test_start_screen_mockup(fake_api, monkeypatch) -> None:
    """8.4 start: live camera component (not st.camera_input / st.file_uploader), «Гараар бичих»,
    «Дугааргүй машин», no «Гарах» here (it stays on home)."""
    monkeypatch.setenv("UI_PASSWORD", "demo-pass")
    at = run_app()
    at.text_input[0].set_value("demo-pass")
    at = click(at, "Нэвтрэх")
    assert "Гарах" in [b.label for b in at.button]                        # home keeps it
    at = on(click(at, "Үйлчилгээ бүртгэх"), RECORD).run()
    assert not at.exception and step(at) == "capture"
    labels = [b.label for b in at.button]
    assert {"Гараар бичих", "Дугааргүй машин"} <= set(labels) and "Гарах" not in labels
    assert not at.get("camera_input") and not at.get("file_uploader")
    assert "Зураг хадгалагдахгүй." in html_blocks(at)


def test_fleet_number_from_start(fake_api) -> None:
    at = click(run_app(RECORD), "Дугааргүй машин")
    assert step(at) == "result" and at.text_input(key="rec_manual").label == "Парк дугаар"
    assert "4 тоо + 3 үсэг" not in html_blocks(at)


def test_photo_size_line_only_with_diag(fake_api, phone) -> None:
    """?diag=1 on ANY page (e.g. home, where login lands) is remembered for the session and the
    record result shows what reached Python (8.4 phone check); hidden otherwise."""
    fake_api.recognize = (200, RECOGNIZED)
    at = phone(run_app(RECORD))
    assert "Зураг: 1920×2560" not in html_blocks(at)
    at = run_app()
    at.query_params["diag"] = "1"
    at = on(click(at.run(), "Үйлчилгээ бүртгэх"), RECORD).run()
    at.query_params.clear()                                   # page switch drops the URL param
    at = phone(at)
    assert "Зураг: 1920×2560 (камер 1920×2560, live, 0 KB)" in html_blocks(at)


def test_iphone_heic_upload_is_sent(fake_api, phone) -> None:
    """RISK-17: a gallery photo the browser could not decode (upper-case .HEIC, no type) is
    sent unchanged to the API as image/heic in the "file" field."""
    seed_existing(fake_api)
    fake_api.recognize = (200, RECOGNIZED)
    at = phone(run_app(RECORD), b"fake-heic", "IMG_0001.HEIC", "", source="gallery")
    assert not at.exception
    body = fake_api.uploads[0]
    assert b'name="file"; filename="IMG_0001.HEIC"' in body
    assert b"Content-Type: image/heic" in body
    at = click(at, "Үргэлжлүүлэх")
    assert step(at) == "form"
    assert "Автоматаар уншсан" in html_blocks(at)                   # strip without the undecodable photo


@pytest.mark.parametrize("kind", ["heic", "jpeg"])
def test_real_format_gallery_photo_sent_unchanged(fake_api, phone, kind) -> None:
    """8.4b: a genuine HEIC (browser could not decode it) and a JPEG from «Галерей» reach the API
    byte for byte with the right content type, and the result screen opens."""
    pillow_heif = pytest.importorskip("pillow_heif")
    import io
    from PIL import Image
    pillow_heif.register_heif_opener()
    buf = io.BytesIO()
    Image.new("RGB", (640, 480), (90, 110, 130)).save(buf, format="HEIF" if kind == "heic" else "JPEG")
    name, ctype = ("IMG_0001.HEIC", "") if kind == "heic" else ("photo.jpg", "image/jpeg")
    fake_api.recognize = (200, RECOGNIZED)
    at = phone(run_app(RECORD), buf.getvalue(), name, ctype, source="gallery")
    assert not at.exception and step(at) == "result"
    body = fake_api.uploads[0]
    assert buf.getvalue() in body
    assert (b"Content-Type: image/heic" if kind == "heic" else b"Content-Type: image/jpeg") in body


def test_browser_read_failure_opens_result_with_diag_error(fake_api, phone) -> None:
    """8.4b: the browser could not read the photo -> result screen (orange, empty field, type the
    plate), nothing sent to the API; the real JS error shows only with ?diag=1."""
    error = "NotReadableError: The requested file could not be read"
    at = phone(run_app(RECORD), error=error, source="gallery")
    assert step(at) == "result" and at.text_input(key="rec_manual").value == ""
    page = html_blocks(at)
    assert "Дугаарыг уншиж чадсангүй — гараар оруулна уу" in page and error not in page
    assert fake_api.uploads == []
    at = run_app()
    at.query_params["diag"] = "1"
    at = on(click(at.run(), "Үйлчилгээ бүртгэх"), RECORD).run()
    at.query_params.clear()
    at = phone(at, error=error, source="gallery")
    assert f"Алдаа: {error} (gallery)" in html_blocks(at)
    at = phone(click(at, "Дахин зураг авах"), tiny_jpeg())       # a good photo next: error line gone
    assert "Алдаа:" not in html_blocks(at)


def test_recognizer_down_falls_back_to_manual(fake_api, phone) -> None:
    at = phone(run_app(RECORD))  # default fake answer: 503
    assert not at.exception
    assert step(at) == "result" and at.text_input(key="rec_manual").value == ""
    assert "Дугаарыг уншиж чадсангүй — гараар оруулна уу" in html_blocks(at)


# --- search + card (5.3 / DESIGN v2) ----------------------------------------------------------

def search(at: AppTest, query: str) -> AppTest:
    at.text_input(key="search_q").set_value(query).run()
    return click(at, "Хайх")


def test_single_match_opens_card(fake_api) -> None:
    seed_existing(fake_api)
    at = search(run_app(SEARCH), "0002 тес")
    assert not at.exception
    assert texts(at.title) == ["Машины карт"]
    page = html_blocks(at)
    assert '<span class="plate lg">0002 ТЕС</span>' in page
    assert "Дараагийн тос солилт" in page and "215 000 км" in page  # 9.5a: target = 200 000 + 15 000
    assert "Сүүлд: 200 000 км · 2026-09-01 · +15 000 км" in page
    assert "Таамаглал" not in page                              # 9.5a: no prediction block
    assert "Шинэ · 1 удаа</span>" in page                       # grey segment chip next to the plate (T7)
    assert "Говь Тест ХХК" in page                             # one line company · model
    assert "Сүүлд солисон" in page and "10W-40 · ердийн\u00a032\u00a0л" in page  # usual liters back (T8)
    assert "Жолооч Д. · 9911 0002" in page and "Залгах" not in page  # driver row, no call button (T6)
    assert texts(at.subheader) == ["Түүх"] and '<div class="tl">' in page  # history timeline, open (T9)


def test_card_shows_target_not_predicted_date(fake_api) -> None:
    """9.5a: even with a prediction, the card shows the target km, not the predicted date."""
    seed_existing(fake_api)
    fake_api.prediction = PREDICTION
    at = search(run_app(SEARCH), "0002 ТЕС")
    [card] = [t for t in texts(at.markdown) if 'class="pc"' in t]
    assert "Дараагийн тос солилт" in card and "215 000 км" in card
    assert "9 хоног" not in card and "10-р сарын 21" not in card


def test_card_mine_truck_has_no_target(fake_api) -> None:
    """9.5a: engine hours only -> no target block; destination row hidden when the service has none."""
    fake_api.add_vehicle("HT-12", "Уурхай А ХХК", "fleet_no", oil_grade="15W-40", liters=40.0, engine_hours=1000.0)
    at = search(run_app(SEARCH), "HT-12")
    page = html_blocks(at)
    assert texts(at.title) == ["Машины карт"] and "1 000 мц" in page
    assert "Дараагийн тос солилт" not in page and 'class="pc"' not in page
    assert "Чиглэл" not in page


def test_search_by_phone_and_recent(fake_api) -> None:
    seed_existing(fake_api)
    at = search(run_app(SEARCH), "+976 9911-0002")
    assert texts(at.title) == ["Машины карт"]
    at = click(on(at, CARD), "Буцах")
    assert texts(at.title) == ["Хайх"]
    assert "Сүүлд үзсэн" in texts(at.subheader)
    assert "0002 ТЕС · Говь Тест ХХК" in [b.label for b in at.button]


def test_several_matches_listed(fake_api) -> None:
    fake_api.add_vehicle("HT-12", "Уурхай А ХХК", "fleet_no", oil_grade="15W-40", liters=40.0, engine_hours=1000.0)
    fake_api.add_vehicle("HT-12", "Уурхай Б ХХК", "fleet_no", oil_grade="10W-30", liters=45.0, engine_hours=2000.0)
    at = search(run_app(SEARCH), "HT-12")
    assert texts(at.title) == ["Хайх"]
    labels = [b.label for b in at.button]
    assert "HT-12 · Уурхай А ХХК · Шинэ" in labels and "HT-12 · Уурхай Б ХХК · Шинэ" in labels


def test_unknown_offers_new_vehicle_form(fake_api) -> None:
    at = search(run_app(SEARCH), "0404 ТЕС")
    assert "Олдсонгүй" in html_blocks(at)
    at = click(at, "Шинэ машин бүртгэх")
    assert texts(at.title) == ["Үйлчилгээ бүртгэх"]
    assert step(at) == "form"
    assert "Шинэ машин" in html_blocks(at)


def test_unknown_phone_message_and_no_new_vehicle_button(fake_api) -> None:
    at = search(run_app(SEARCH), "99119999")
    assert "Энэ утасны дугаартай үйлчилгээ бүртгэгдээгүй" in html_blocks(at)
    assert "Шинэ машин бүртгэх" not in [b.label for b in at.button]  # a phone is not a plate


def test_search_by_camera(fake_api) -> None:
    seed_existing(fake_api)
    fake_api.recognize = (200, RECOGNIZED)
    at = click(run_app(SEARCH), "Камераар хайх")
    at = upload_photo(at, "search_upload")
    assert texts(at.title) == ["Машины карт"]


def test_card_record_button_opens_prefilled_form(fake_api) -> None:
    seed_existing(fake_api)
    at = click(on(search(run_app(SEARCH), "0002 ТЕС"), CARD), "Үйлчилгээ бүртгэх")
    assert texts(at.title) == ["Үйлчилгээ бүртгэх"]
    assert step(at) == "form"
    assert at.session_state["f_oil_choice"] == "10W-40"


def test_save_then_find_by_plate_and_phone(fake_api) -> None:
    """DONE-WHEN path: save a new truck, then find it by plate AND by phone."""
    at = click(fill_new_vehicle(to_form(run_app(RECORD), "0001 ТЕС")), "Хадгалах")
    at = click(at, "Машины карт үзэх")
    assert texts(at.title) == ["Машины карт"]
    assert '<span class="plate lg">0001 ТЕС</span>' in html_blocks(at)

    at = search(run_app(SEARCH), "0001 ТЕС")
    assert '<span class="plate lg">0001 ТЕС</span>' in html_blocks(at)
    at = search(run_app(SEARCH), "99110099")
    assert '<span class="plate lg">0001 ТЕС</span>' in html_blocks(at)
    assert "Тодорхойгүй" in html_blocks(at)  # new plate vehicle saved with the default company (T8)


# --- manager dashboard (5.4) ---------------------------------------------------------------

def reminder(service_id: int, plate: str, days_left: int, status: str) -> dict:
    return {
        "service_id": service_id, "vehicle_id": service_id, "identifier_type": "plate",
        "identifier": plate, "plate_normalized": plate.replace(" ", ""), "company": "Туршилт ХХК",
        "phone": "99110002", "driver_name": "Тест Д.", "last_change_date": "2026-08-10",
        "oil_grade": "15W-40", "liters": 32.0, "method": "b0", "buffer_days": 12,
        "predicted_date": "2026-10-07", "send_on": "2026-09-25", "days_left": days_left, "status": status,
        # like the API (contract v0.7): at_risk = predicted date + 7 days grace has passed
        "segment": "at_risk" if days_left < -7 else "loyal",
        "message_mn": f"Сайн байна уу! {plate} дугаартай машины тос солих хугацаа болсон байна.",
    }


@pytest.mark.parametrize(("key", "title"), [
    ("home_record", "Үйлчилгээ бүртгэх"), ("home_search", "Хайх"), ("home_reminders", "Сануулга"),
    ("home_dashboard", "Хяналтын самбар")])
def test_home_card_navigates(fake_api, key, title) -> None:
    """5.9b: every home card is a real button that opens its screen (the «Сануулга» card was a
    page_link whose <a> covered only one text line of the card, so taps did nothing)."""
    at = run_app()
    at.button(key=key).click().run()
    assert texts(at.title) == [title]


def test_home_dashboard_button(fake_api) -> None:
    """8.1: «Хяналтын самбар» on home opens the manager screen (5.4) with its segment tiles."""
    at = run_app()
    assert at.button(key="home_dashboard").label == "Хяналтын самбар"
    at.button(key="home_dashboard").click().run()
    assert not at.exception
    assert texts(at.subheader) == ["Сегмент", "Машин хайх"]
    assert '<div class="tiles two seg">' in html_blocks(at)


def test_css_stops_iphone_tap_zoom() -> None:
    """8.1: fields are 16 px (iOS zooms on focus below 16 px) and tap targets have
    touch-action: manipulation; pinch zoom is never switched off."""
    from src.ui import common
    css = common.CSS
    assert ".stApp input, .stApp textarea, .stApp select, [data-baseweb=\"select\"] > div {font-size: 16px !important;}" in css
    assert ".stApp button, .stApp a" in css and "touch-action: manipulation;" in css
    assert "maximum-scale" not in css and "user-scalable" not in css


def test_home_reminder_summary(fake_api) -> None:
    """T9: greeting line, 3 numbers and the red badge from GET /reminders/due (+ total from /stats/segments)."""
    fake_api.reminders = [reminder(1, "0001 ТЕС", 5, "due_soon"), reminder(3, "0003 ТЕС", -20, "overdue")]
    at = run_app()
    page = html_blocks(at)
    assert "Өнөөдөр 1 машин сануулга хүлээж байна." in page
    assert 'Сануулах</div><div class="v">1</div>' in page
    assert 'Хэтэрсэн</div><div class="v" style="color:#B42318">1</div>' in page
    assert 'Машин</div><div class="v">7</div>' in page
    assert '<span class="badge">1</span>' in page


def test_home_without_predictions_still_works(fake_api) -> None:
    fake_api.reminders = None  # -> 503 (model not loaded)
    at = run_app()
    assert not at.exception and not at.error
    assert "Тос солилтын бүртгэл, сануулга." in html_blocks(at)
    assert '<span class="badge">' not in html_blocks(at)


def test_dashboard_tiles_and_two_groups(fake_api) -> None:
    fake_api.reminders = [reminder(1, "0001 ТЕС", 5, "due_soon"), reminder(2, "0002 ТЕС", -3, "overdue"),
                          reminder(3, "0003 ТЕС", -20, "overdue")]
    at = run_app(REMINDERS)
    assert not at.exception
    page = html_blocks(at)
    for label, count in (("Өнөөдөр", 2), ("Эрсдэлтэй", 1)):
        assert f'<div class="v">{count}</div><div class="k">{label}</div>' in page
    assert '<div class="v" style="color:#B42318">2</div><div class="k">Хэтэрсэн</div>' in page
    assert texts(at.subheader) == ["Өнөөдрийн дараалал (2)", "Эрсдэлтэй (1)"]
    assert "0003 ТЕС" in page and "Тест Д." in page and "9911\u00a00002" in page
    assert links(at)["Залгах"].url == "tel:99110002"
    assert [c.value for c in at.code][0].startswith("Сайн байна уу! 0001 ТЕС")  # copyable text
    assert "GEO_EXPERT_PHONE" in page and not at.warning  # placeholder contact = one small line (T7)


def links(at: AppTest) -> dict[str, Any]:
    """st.link_button elements by label (AppTest can read them but not click them)."""
    return {b.proto.label: b.proto for b in at.get("link_button")}


def test_viber_button_opens_prefilled_message(fake_api) -> None:
    """5.5 (D-091): one Viber link per reminder, text = the exact message_mn, and the tap is a
    widget (has an id) so the app reruns and can ask «Илгээсэн үү?». Copy box stays (fallback)."""
    fake_api.reminders = [reminder(1, "0001 ТЕС", 5, "due_soon"), reminder(3, "0003 ТЕС", -20, "overdue")]
    at = run_app(REMINDERS)
    vibers = [b.proto for b in at.get("link_button") if b.proto.label == "Viber-ээр илгээх"]
    assert len(vibers) == 2
    url = vibers[0].url
    assert url.startswith("viber://forward?text=") and " " not in url and "ТЕС" not in url
    assert unquote(url.removeprefix("viber://forward?text=")) == reminder(1, "0001 ТЕС", 5, "due_soon")["message_mn"]
    assert vibers[0].id.endswith("viber_1") and not links(at)["Залгах"].id  # tap reruns; tel: does not
    assert [c.value for c in at.code][0].startswith("Сайн байна уу! 0001 ТЕС")
    assert "Илгээсэн үү?" not in html_blocks(at)  # asked only after the tap


def tapped_viber(service_id: int) -> AppTest:
    """Reminders screen as it is after the Viber link of this row was tapped (callback state)."""
    at = run_app(REMINDERS)
    at.session_state["rem_viber_tapped"] = {service_id}
    return at.run()


def test_viber_tap_asks_then_yes_marks_sent(fake_api) -> None:
    fake_api.reminders = [reminder(1, "0001 ТЕС", 5, "due_soon"), reminder(2, "0002 ТЕС", -3, "overdue")]
    at = tapped_viber(1)
    assert "Илгээсэн үү?" in html_blocks(at)
    assert [b.label for b in at.button].count("Илгээсэн") == 1  # row 2 keeps the plain button
    at = click(at, "Тийм")
    assert any("Илгээсэн гэж тэмдэглэлээ" in t for t in texts(at.success))
    assert list(fake_api.sent) == [1]                             # POST /reminders/1/sent
    assert "Илгээсэн үү?" not in html_blocks(at)


def test_viber_tap_no_hides_question(fake_api) -> None:
    fake_api.reminders = [reminder(1, "0001 ТЕС", 5, "due_soon")]
    at = click(tapped_viber(1), "Үгүй")
    assert "Илгээсэн үү?" not in html_blocks(at) and "Илгээсэн" in [b.label for b in at.button]
    assert not fake_api.sent


def test_no_viber_button_without_message(fake_api) -> None:
    fake_api.reminders = [{**reminder(1, "0001 ТЕС", 5, "due_soon"), "message_mn": ""}]
    at = run_app(REMINDERS)
    assert "Viber-ээр илгээх" not in links(at) and "Илгээсэн" in [b.label for b in at.button]


def test_mark_sent_removes_row(fake_api) -> None:
    fake_api.reminders = [reminder(1, "0001 ТЕС", 5, "due_soon")]
    at = click(run_app(REMINDERS), "Илгээсэн")
    assert any("Илгээсэн гэж тэмдэглэлээ" in t for t in texts(at.success))
    assert texts(at.subheader)[0] == "Өнөөдрийн дараалал (0)"  # gone from the due list


def test_mark_sent_twice_shows_already_sent(fake_api) -> None:
    fake_api.reminders = [reminder(1, "0001 ТЕС", 5, "due_soon")]
    fake_api.hide_sent = False  # the row stays, as in a second manager's stale tab
    at = click(click(run_app(REMINDERS), "Илгээсэн"), "Илгээсэн")
    assert any("Өмнө нь илгээсэн" in t for t in texts(at.info))
    assert len(fake_api.sent) == 1  # stored once (API is idempotent)


def test_mark_sent_not_latest_service(fake_api) -> None:
    fake_api.reminders = [reminder(999, "0009 ТЕС", 5, "due_soon")]
    at = click(run_app(REMINDERS), "Илгээсэн")
    assert any("шинэ үйлчилгээтэй" in t for t in texts(at.warning))


def test_overview_segment_tiles(fake_api) -> None:
    at = click(run_app(REMINDERS), "Машинуудын тойм харах")
    assert texts(at.title) == ["Хяналтын самбар"]
    page = html_blocks(at)
    for label, count in (("Шинэ", 2), ("Давтан", 3), ("Байнгын", 1), ("Эрсдэлтэй", 1)):
        assert f'<div class="v">{count}</div><div class="k">' in page and f"</i>{label}</div>" in page
    assert "Нийт 7 машин" in page


def test_dashboard_without_model(fake_api) -> None:
    fake_api.reminders = None
    at = run_app(REMINDERS)
    assert not at.exception
    assert any("загвар ачаалагдаагүй" in t for t in texts(at.info))


# --- «Хяналтын самбар» search + history (9.1) -------------------------------------------------

def seed_two(fake: FakeApi) -> None:
    seed_existing(fake)  # 0002 ТЕС · Говь Тест ХХК
    fake.add_vehicle("0003 ТЕС", "Хан Тест ХХК", service_date="2026-09-05", odometer_km=150000,
                     oil_grade="15W-40", liters=30.0, destination_id=2, driver_name="Бат Т.", driver_phone="99110003")


def test_dashboard_lists_every_truck_then_searches(fake_api) -> None:
    seed_two(fake_api)
    at = run_app(DASHBOARD)
    assert not at.exception
    assert fake_api.searches == [""]                         # blank search = every truck
    assert len(at.dataframe[0].value) == 2
    assert "Бүх машин · 2" in html_blocks(at)
    at.text_input(key="dash_q").set_value("говь")
    at = click(on(at, DASHBOARD), "Хайх")
    assert fake_api.searches[-1] == "говь"
    table = at.dataframe[0].value
    assert list(table["Дугаар"]) == ["0002 ТЕС"]
    assert list(table.columns) == ["Дугаар", "Компани", "Сүүлд солисон", "Гүйлт (км)", "Дараагийн солилт (км)",
                                   "Тос", "Литр", "Жолооч", "Утас", "Удаа"]  # 9.5a: target instead of «Төлөв»
    assert list(table["Дараагийн солилт (км)"]) == [215000]
    assert at.dataframe[0].proto.placeholder == "—"  # 9.5a: empty cells show «—», never "None"
    assert "«говь» · 1 машин" in html_blocks(at)


def test_dashboard_nothing_found(fake_api) -> None:
    seed_two(fake_api)
    at = run_app(DASHBOARD)
    at.text_input(key="dash_q").set_value("ЯЯЯ")
    at = click(on(at, DASHBOARD), "Хайх")
    assert not at.dataframe and "Олдсонгүй." in html_blocks(at)


def test_dashboard_phone_card_opens_full_history(fake_api) -> None:
    """Phone view: tapping a truck's row shows every service: date, odometer, oil, liters, driver, phone."""
    seed_two(fake_api)
    fake_api._add_service(1, {"service_date": "2026-10-01", "odometer_km": 215000, "oil_grade": "15W-40",
                              "liters": 30.0, "driver_name": "Шинэ Ж.", "driver_phone": "99110009"})
    at = run_app(DASHBOARD)
    at = on(at, DASHBOARD).button(key="dash_open_1").click().run()
    page = html_blocks(at)
    assert '<table class="hist">' in page
    for th in ("Огноо", "Гүйлт → дараагийн", "Тос", "Жолооч"):  # liters in «Тос», phone under the driver
        assert f"<th>{th}</th>" in page
    for target in ("230 000 км", "215 000 км"):  # 9.5a: each service's own target, under its odometer
        assert f"<div class=ph>→ {target}</div>" in page
    assert page.index("2026-10-01") < page.index("2026-09-01")   # newest first
    for value in ("215 000 км", "200 000 км", "10W-40 · 32 л", "15W-40 · 30 л", "Жолооч Д.", "Шинэ Ж.",
                  '<div class="ph">9911 0002</div>', '<div class="ph">9911 0009</div>'):
        assert value in page
    at = on(at, DASHBOARD).button(key="dash_close").click().run()
    assert '<table class="hist">' not in html_blocks(at) and len(at.dataframe[0].value) == 2  # back to the list


def test_dashboard_table_row_opens_history(fake_api) -> None:
    """Computer view: a table click calls on_table_select (row number -> vehicle id)."""
    seed_two(fake_api)
    at = run_app(DASHBOARD)
    assert at.session_state["dash_ids"] == [1, 2]
    at.session_state["dash_vid"] = 2   # what on_table_select stores for row 1 (AppTest cannot click a grid row)
    at = on(at, DASHBOARD).run()
    assert "Бат Т." in html_blocks(at) and '<table class="hist">' in html_blocks(at)


def test_dashboard_due_list_is_read_only(fake_api) -> None:
    fake_api.reminders = [reminder(1, "0001 ТЕС", 5, "due_soon"), reminder(3, "0003 ТЕС", -20, "overdue")]
    at = run_app(DASHBOARD)
    page = html_blocks(at)
    assert "Сануулах машин (2)" in page and "0001 ТЕС" in page and "0003 ТЕС" in page
    assert not {"Илгээсэн", "Тийм"} & {b.label for b in at.button} and not at.get("link_button")


def test_dashboard_writes_nothing(fake_api) -> None:
    """View only (9.1): opening, searching and opening a history never POSTs."""
    seed_two(fake_api)
    fake_api.reminders = [reminder(1, "0001 ТЕС", 5, "due_soon")]
    posts_before = len(fake_api.posts)
    at = run_app(DASHBOARD)
    at.text_input(key="dash_q").set_value("ТЕС")
    at = click(on(at, DASHBOARD), "Хайх")
    on(at, DASHBOARD).button(key="dash_open_2").click().run()
    assert len(fake_api.posts) == posts_before and not fake_api.sent and not fake_api.uploads


def test_wide_layout_only_on_dashboard(fake_api) -> None:
    """9.1: ≥ 1024 px the dashboard uses the full width; every other screen keeps the 480 px column."""
    from src.ui import common
    assert ".block-container {padding: 12px 20px 40px !important; max-width: 480px;}" in common.CSS
    at = run_app(DASHBOARD)
    assert "@media (min-width: 1024px)" in html_blocks(at) and "max-width: 1360px" in html_blocks(at)
    for view in (HOME, SEARCH, CARD, RECORD, REMINDERS):
        assert "max-width: 1360px" not in html_blocks(run_app(view)), view


def set_filters(at: AppTest, **values) -> AppTest:
    """Fill dashboard filters (widget keys dash_f_<param>) and press «Хайх»."""
    at = on(at, DASHBOARD)
    for name, value in values.items():
        widget = at.date_input(key=f"dash_f_{name}") if name.startswith("date") else at.selectbox(key=f"dash_f_{name}")
        widget.set_value(value)
    return click(at, "Хайх")


def test_dashboard_filters_sent_with_search(fake_api) -> None:
    from datetime import date
    seed_two(fake_api)
    at = run_app(DASHBOARD)
    assert at.selectbox(key="dash_f_company_id").options == ["Говь Тест ХХК", "Хан Тест ХХК"]
    assert at.selectbox(key="dash_f_oil_grade").options == ["10W-40", "15W-40"]
    assert at.selectbox(key="dash_f_route").options[1] == "Монгол – Хятад"
    at = set_filters(at, date_from=date(2026, 9, 1), date_to=date(2026, 9, 30), company_id=2,
                     route="mn_china", oil_grade="15W-40")
    assert not at.exception
    sent = fake_api.search_params[-1]
    assert sent == {"q": "", "limit": "1000", "with_services": "true", "date_from": "2026-09-01", "date_to": "2026-09-30",
                    "company_id": "2", "route": "mn_china", "oil_grade": "15W-40"}
    assert list(at.dataframe[0].value["Дугаар"]) == ["0003 ТЕС"]
    assert at.expander[0].label == "Шүүлтүүр (5)"


def test_dashboard_filter_dates_in_wrong_order(fake_api) -> None:
    from datetime import date
    seed_two(fake_api)
    at = run_app(DASHBOARD)
    calls = len(fake_api.search_params)
    at = set_filters(at, date_from=date(2026, 10, 1), date_to=date(2026, 9, 1))
    assert any("хойш байна" in w for w in texts(at.warning))
    assert len(fake_api.search_params) == calls          # not sent (API would answer 422)
    assert at.selectbox(key="dash_f_company_id").options  # lists kept from the last good answer


def test_dashboard_clear_filters(fake_api) -> None:
    seed_two(fake_api)
    at = set_filters(run_app(DASHBOARD), company_id=1)
    assert len(at.dataframe[0].value) == 1
    at = on(at, DASHBOARD).button(key="dash_clear").click().run()
    assert fake_api.search_params[-1] == {"q": "", "limit": "1000", "with_services": "true"}
    assert len(at.dataframe[0].value) == 2 and at.expander[0].label == "Шүүлтүүр"


def test_dashboard_excel_has_filtered_list(fake_api, monkeypatch) -> None:
    """«Excel татах»: one row per service of the filtered trucks + a line saying which filters were used."""
    import io
    from openpyxl import load_workbook
    from src.ui import export
    made: list[bytes] = []
    real = export.search_workbook

    def spy(vehicles, today, filters_text):
        made.append(real(vehicles, today, filters_text))
        return made[-1]

    monkeypatch.setattr(export, "search_workbook", spy)
    seed_two(fake_api)
    fake_api._add_service(1, {"service_date": "2026-10-01", "odometer_km": 215000, "oil_grade": "15W-40",
                              "liters": 30.0, "driver_name": "Шинэ Ж.", "driver_phone": "99110009"})
    at = set_filters(run_app(DASHBOARD), company_id=1)
    assert len(at.dataframe[0].value) == 1                         # the screen: one row per truck
    [button] = [b.proto for b in at.get("download_button")]
    assert button.label == "Excel татах" and button.ignore_rerun   # downloading does not rerun the page
    ws = load_workbook(io.BytesIO(made[-1])).active
    assert ws["A2"].value == "Компани: Говь Тест ХХК"
    rows = [[c.value for c in r][:3] for r in ws.iter_rows(min_row=5) if r[0].value]
    assert [(r[0].date().isoformat(), r[1], r[2]) for r in rows] == [   # the file: one row per service
        ("2026-10-01", "0002 ТЕС", "Говь Тест ХХК"), ("2026-09-01", "0002 ТЕС", "Говь Тест ХХК")]


def test_dashboard_no_excel_button_when_nothing_found(fake_api) -> None:
    at = run_app(DASHBOARD)  # empty database
    assert not [b for b in at.get("download_button")]


def test_dashboard_table_toolbar_hidden(fake_api) -> None:
    """9.1 (Ari): no CSV download/search toolbar on the table; export only via the Excel button."""
    seed_two(fake_api)
    assert ('.st-key-dash_table_wrap [data-testid="stElementToolbar"] {display: none !important;}'
            in html_blocks(run_app(DASHBOARD)))
