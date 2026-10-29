"""Tests for src/ui/export.py: the dashboard's Excel file, one row per service (task 9.1 step 2).
Invented data only (RISK-06)."""

import io
import zipfile
from datetime import date

import pytest
from openpyxl import load_workbook

from src.ui import export

S1 = {"service_date": "2026-09-01", "odometer_km": 205000, "target_km": 220000, "engine_hours": None, "oil_grade": "15W-40",
      "liters": 30.0, "route": "mn_china", "driver_name": "Тест Д.", "driver_phone": "99110001"}
S2 = {**S1, "service_date": "2026-06-10", "odometer_km": 182000, "target_km": 197000, "route": "mn_russia", "driver_name": "Бат Т."}
ROW = {"id": 1, "identifier_type": "plate", "identifier": "0001 ТЕС", "vehicle_model": None,
       "company": {"id": 1, "name": "Туршилт ХХК"}, "service_count": 2, "last_service": S1,
       "segment": "returning", "status": None, "days_left": None, "matched_on": [], "services": [S1, S2]}
OTHER = {**ROW, "id": 2, "identifier": "0002 ТЕС", "company": {"id": 2, "name": "Хоёр ХХК"},
         "services": [{**S1, "service_date": "2026-07-01", "route": None, "odometer_km": None, "target_km": None,
                       "driver_phone": None}]}
TODAY = date(2026, 10, 28)


def book(vehicles: list[dict], text: str = "Шүүлтүүргүй — бүх машин") -> bytes:
    return export.search_workbook(vehicles, TODAY, text)


def sheet(vehicles: list[dict], text: str = "Шүүлтүүргүй — бүх машин"):
    return load_workbook(io.BytesIO(book(vehicles, text))).active


def data_rows(ws) -> list[list]:
    return [[c.value for c in r] for r in ws.iter_rows(min_row=export.HEADER_ROW + 1) if r[0].value is not None]


def test_one_row_per_service_newest_first() -> None:
    ws = sheet([ROW, OTHER])
    assert ws.title == "Үйлчилгээ"
    assert ws["A1"].value == "Гео-Эксперт · үйлчилгээ · 2026-10-28"
    assert [c.value for c in ws[export.HEADER_ROW]] == [
        "Огноо", "Дугаар", "Компани", "Маршрут", "Тос", "Литр", "Гүйлт (км)", "Моторын цаг",
        "Дараагийн солилт (км)", "Жолооч", "Утас"]
    rows = data_rows(ws)
    assert [(r[0].date().isoformat(), r[1]) for r in rows] == [
        ("2026-09-01", "0001 ТЕС"), ("2026-07-01", "0002 ТЕС"), ("2026-06-10", "0001 ТЕС")]
    assert rows[0][1:] == ["0001 ТЕС", "Туршилт ХХК", "Монгол – Хятад", "15W-40", 30.0, 205000, None,
                           220000, "Тест Д.", "9911 0001"]  # 9.5a: target km after the meters
    assert rows[2][8] == 197000  # each service has its own target
    assert rows[1][3] is None and rows[1][6] is None and rows[1][8] is None and rows[1][10] is None  # missing -> empty
    assert rows[2][3] == "Монгол – Орос"


def test_header_bold_frozen_with_excel_filter() -> None:
    ws = sheet([ROW])
    assert ws.cell(export.HEADER_ROW, 1).font.bold
    assert ws.freeze_panes == "A5"
    assert ws.auto_filter.ref == "A4:K6"


@pytest.mark.parametrize("evil", ["=1+1", '=HYPERLINK("http://example.test","x")', "=cmd|' /C calc'!A0"])
def test_text_starting_with_equals_is_never_a_formula(evil: str) -> None:
    """Typed values (driver, company, plate, filter line) are written as plain text, not formulas."""
    service = {**S1, "driver_name": evil, "oil_grade": evil}
    vehicle = {**ROW, "identifier": evil, "company": {"id": 9, "name": evil}, "services": [service]}
    data = book([vehicle], evil)
    ws = load_workbook(io.BytesIO(data)).active
    for cell in (ws["A2"], ws["B5"], ws["C5"], ws["E5"], ws["J5"]):  # filter line, plate, company, oil, driver
        assert cell.value == evil and cell.data_type == "s" and cell.number_format == "@"
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert "<f>" not in z.read("xl/worksheets/sheet1.xml").decode()  # no formula anywhere in the sheet


def test_numbers_and_dates_stay_numbers_and_dates() -> None:
    ws = sheet([ROW])
    assert ws["A5"].data_type == "d" and ws["F5"].data_type == "n" and ws["G5"].data_type == "n"
    assert ws["I5"].value == 220000 and ws["I5"].data_type == "n" and ws["I5"].number_format == "#,##0"  # 9.5a


def test_engine_hours_next_to_odometer_for_mine_trucks() -> None:
    """9.1c: a mine truck has engine hours instead of km (D-006): hours filled, odometer empty."""
    mine = {**S1, "route": "mine_site", "odometer_km": None, "target_km": None, "engine_hours": 12980.5}
    ws = sheet([{**ROW, "identifier_type": "fleet_no", "identifier": "HT-12", "services": [mine]}])
    header = [c.value for c in ws[export.HEADER_ROW]]
    assert header.index("Моторын цаг") == header.index("Гүйлт (км)") + 1
    assert (ws["G5"].value, ws["H5"].value) == (None, 12980.5)
    assert ws["H5"].data_type == "n" and ws["H5"].number_format == "#,##0.0"
    assert ws["D5"].value == "Уурхай"
    assert ws["I5"].value is None  # 9.5a: no target km for a mine truck


def test_empty_list_still_valid_file() -> None:
    ws = sheet([])
    assert ws.cell(export.HEADER_ROW, 1).value == "Огноо" and ws.cell(export.HEADER_ROW + 1, 1).value is None


def test_file_name_is_ascii() -> None:
    name = export.file_name(TODAY)
    assert name == "geo_expert_uilchilgee_2026-10-28.xlsx" and name.isascii()
