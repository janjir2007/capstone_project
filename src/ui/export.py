"""Excel export of the dashboard's filtered list (task 9.1 step 2): ONE ROW PER MATCHING SERVICE.

Built in memory and sent only to the manager's browser (st.download_button); nothing is stored
on the server. The file holds driver names and phones (Geo-Expert's own customer data), so it is
for internal use only (RISK-06). No Streamlit here, so it is tested on its own.

Formula safety: names, companies and phones are typed by people. A text cell starting with "="
would become an Excel formula, so every text cell is written as plain text (type "s" + format "@").
"""

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from src.ui import logic

SHEET = "Үйлчилгээ"
HEADER_ROW = 4  # rows 1–2: title + filters, row 3 empty
DATE_COLUMN = "Огноо"
WIDTHS = {"Огноо": 12, "Дугаар": 12, "Компани": 26, "Маршрут": 18, "Тос": 10, "Литр": 7,
          "Гүйлт (км)": 12, "Жолооч": 18, "Утас": 12}


def file_name(today: date) -> str:
    """ASCII name (some phones mangle Cyrillic file names): geo_expert_uilchilgee_2026-10-28.xlsx."""
    return f"geo_expert_uilchilgee_{today.isoformat()}.xlsx"


def put(ws, row: int, col: int, value) -> Cell:
    """Write one cell; text is ALWAYS plain text, never a formula ("=SUM(..)" stays those characters)."""
    cell = ws.cell(row, col)
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"     # openpyxl made it "f" (formula) when the text starts with "="
        cell.number_format = "@"  # Excel's Text format, also when the cell is edited later
    return cell


def search_workbook(vehicles: list[dict], today: date, filters_text: str) -> bytes:
    """.xlsx bytes: title, the filters used, then one row per service (logic.EXPORT_COLUMNS).

    ``vehicles`` = GET /vehicles/search?with_services=true rows. Dates are real Excel dates and
    numbers real numbers; the header row is bold, frozen and has Excel's own filter buttons.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    put(ws, 1, 1, f"Гео-Эксперт · үйлчилгээ · {today.isoformat()}").font = Font(bold=True, size=13)
    put(ws, 2, 1, filters_text)

    columns = logic.EXPORT_COLUMNS
    rows = logic.service_export_rows(vehicles)
    for j, name in enumerate(columns, start=1):
        put(ws, HEADER_ROW, j, name).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS.get(name, 12)
    for i, row in enumerate(rows, start=HEADER_ROW + 1):
        for j, name in enumerate(columns, start=1):
            value = row[name]
            if name == DATE_COLUMN and value:
                cell = put(ws, i, j, date.fromisoformat(value))
                cell.number_format = "yyyy-mm-dd"
            else:
                cell = put(ws, i, j, value)
                if name == "Гүйлт (км)":
                    cell.number_format = "#,##0"
    ws.auto_filter.ref = f"A{HEADER_ROW}:{get_column_letter(len(columns))}{HEADER_ROW + max(len(rows), 1)}"
    ws.freeze_panes = f"A{HEADER_ROW + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
